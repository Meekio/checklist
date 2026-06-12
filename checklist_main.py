import os
import json
import time
from math import ceil
from datetime import datetime

import pandas as pd
import fitz  # PyMuPDF
from dotenv import load_dotenv
from google import genai
from google.genai import types
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

load_dotenv()

GEMINI_MODEL_NAME = "gemini-2.5-flash"
RULES_PER_BATCH = 20
API_CALL_DELAY_SECONDS = 2
API_TIMEOUT_SECONDS = 120

# This common rules file is always included in the QC process.
COMMON_RULES_FILE_PATH = os.path.join(
    os.path.dirname(os.path.abspath(__file__)), "files","Y14.5_4.1_Fundamental_Rules_Conditional_Checklist.xlsx")

# Gemini helpers

def initialize_gemini_client() -> genai.Client:
    """
    Create and return a Gemini client using the API key from the .env file.
    """
    gemini_api_key = os.getenv("GEMINI_API_KEY")

    if not gemini_api_key:
        raise EnvironmentError(
            "GEMINI_API_KEY was not found in the environment variables or .env file."
        )

    print("[INFO] Gemini client initialized successfully.")
    return genai.Client(api_key=gemini_api_key)


def build_gemini_prompt(rules_batch: list[str]) -> str:
    """
    Build the text prompt sent to Gemini for one batch of QC rules.
    """
    numbered_rules_lines = []

    for rule_number, rule_text in enumerate(rules_batch, start=1):
        numbered_rules_lines.append(f"{rule_number}. {rule_text}")

    numbered_rules_text = "\n".join(numbered_rules_lines)

    prompt = f"""
You are an expert engineering drawing quality inspector.

Evaluate the attached engineering drawing (all pages are provided as images) against the following QC rules.

Return STRICT JSON ONLY.
The response must be a top-level array with one object per rule in the SAME ORDER as the rules given below.

Rules to evaluate:
{numbered_rules_text}

Expected JSON response format:
[
  {{
    "status": "Correct or Incorrect",
    "remarks": "Very short reason"
  }},
  ...
]
""".strip()

    return prompt


def evaluate_rules_batch_with_gemini(
    gemini_client: genai.Client,
    drawing_image_parts: list,
    rules_batch: list[str]
) -> str:
    """
    Send one batch of rules and all drawing pages to Gemini.
    Returns the raw text response from Gemini.
    """
    prompt_text = build_gemini_prompt(rules_batch)
    request_contents = [prompt_text] + drawing_image_parts

    try:
        response = gemini_client.models.generate_content(
            model=GEMINI_MODEL_NAME,
            contents=request_contents,
            config=types.GenerateContentConfig(
                http_options=types.HttpOptions(timeout=API_TIMEOUT_SECONDS * 1000)
            ),
        )
        return response.text

    except Exception as error:
        raise RuntimeError(f"Gemini API call failed: {error}")

# PDF processing

def convert_pdf_to_gemini_image_parts(pdf_file_path: str) -> list:
    """
    Convert each page of the PDF drawing into an image part that can be sent to Gemini.
    """
    if not os.path.isfile(pdf_file_path):
        raise FileNotFoundError(f"Drawing PDF not found: {pdf_file_path}")

    print(f"[INFO] Rendering PDF drawing: {pdf_file_path}")

    pdf_document = fitz.open(pdf_file_path)
    drawing_image_parts = []

    try:
        for page_index in range(len(pdf_document)):
            pdf_page = pdf_document[page_index]
            page_pixmap = pdf_page.get_pixmap(dpi=300)

            image_part = types.Part.from_bytes(
                data=page_pixmap.tobytes("png"),
                mime_type="image/png"
            )

            drawing_image_parts.append(image_part)
            print(f"[INFO] Rendered page {page_index + 1}")

    finally:
        pdf_document.close()

    if not drawing_image_parts:
        raise ValueError("No pages were rendered from the PDF.")

    print(f"[INFO] Total PDF pages rendered: {len(drawing_image_parts)}")
    return drawing_image_parts

# Excel rules loading

def find_rules_column_name(rules_dataframe: pd.DataFrame) -> str:
    """
    Identify which column in the Excel file contains the QC rule text.

    Priority:
    1. Exact match: 'QC Point'
    2. Known alternative names
    3. First non-empty column
    """
    column_names = list(rules_dataframe.columns)

    normalized_column_names = {
        column_name: str(column_name).strip().lower()
        for column_name in column_names
    }

    if "QC Point" in rules_dataframe.columns:
        return "QC Point"

    possible_rule_column_names = [
        "qc point",
        "qc points",
        "qc rule",
        "qc rules",
        "rule",
        "rules",
        "check",
        "check point",
        "checklist",
        "qc",
        "requirement",
        "criteria",
        "phenomenon / condition",
        "phenomenon/condition",
    ]

    for column_name in column_names:
        normalized_name = normalized_column_names[column_name]
        if normalized_name in possible_rule_column_names:
            return column_name

    for column_name in column_names:
        column_has_any_text = (
            rules_dataframe[column_name]
            .astype(str)
            .str.strip()
            .replace({"nan": ""})
            .str.len()
            .gt(0)
            .any()
        )

        if column_has_any_text:
            print(
                f"[WARN] Could not find a standard rules column name. "
                f"Using column '{column_name}' as fallback. "
                f"Consider renaming it to 'QC Point' for clarity."
            )
            return column_name

    raise ValueError(
        "Could not find any usable rules column in the Excel file. "
        f"Available columns: {', '.join(map(str, column_names))}"
    )


def load_rules_excel_file(excel_file_path: str, rules_label: str) -> tuple[pd.DataFrame, str]:
    """
    Load an Excel file containing QC rules and return:
    - the dataframe
    - the detected rule column name
    """
    print(f"[INFO] Loading {rules_label} rules from: {excel_file_path}")

    try:
        rules_dataframe = pd.read_excel(excel_file_path, engine="openpyxl")
    except Exception:
        rules_dataframe = pd.read_excel(excel_file_path)

    rules_column_name = find_rules_column_name(rules_dataframe)

    print(
        f"[INFO] Loaded {len(rules_dataframe)} rows from {rules_label} rules. "
        f"Detected rule column: '{rules_column_name}'"
    )

    return rules_dataframe, rules_column_name

# Utility helpers

def remove_markdown_code_fences(text: str) -> str:
    """
    Remove markdown code fences such as ```json ... ``` from Gemini output,
    if they exist.
    """
    cleaned_text = text.strip()

    if cleaned_text.startswith("```"):
        cleaned_text = cleaned_text.split("```", 2)[1]

        if cleaned_text.startswith("json"):
            cleaned_text = cleaned_text[4:]

        cleaned_text = cleaned_text.rsplit("```", 1)[0].strip()

    return cleaned_text


def split_list_into_batches(items: list, batch_size: int):
    """
    Yield batches from a list along with the starting index of each batch.
    """
    for batch_start_index in range(0, len(items), batch_size):
        batch_items = items[batch_start_index: batch_start_index + batch_size]
        yield batch_items, batch_start_index

# Rule evaluation

def prepare_non_empty_rules(
    rules_dataframe: pd.DataFrame,
    rules_column_name: str
) -> tuple[list[str], list[int], list[str], list[str]]:
    """
    Extract non-empty rules from the rules dataframe.

    Returns:
    - non_empty_rule_texts
    - original_row_indexes
    - initial_statuses
    - initial_remarks

    Blank rules are marked as Error immediately.
    """
    rules_series = rules_dataframe[rules_column_name]

    non_empty_rule_texts = []
    original_row_indexes = []

    initial_statuses = [""] * len(rules_dataframe)
    initial_remarks = [""] * len(rules_dataframe)

    for row_index, rule_value in enumerate(rules_series):
        if pd.isna(rule_value):
            cleaned_rule_text = ""
        else:
            cleaned_rule_text = str(rule_value).strip()

        if cleaned_rule_text:
            non_empty_rule_texts.append(cleaned_rule_text)
            original_row_indexes.append(row_index)
        else:
            initial_statuses[row_index] = "Error"
            initial_remarks[row_index] = "Blank rule text"

    return non_empty_rule_texts, original_row_indexes, initial_statuses, initial_remarks


def evaluate_rules_against_drawing(
    gemini_client: genai.Client,
    drawing_image_parts: list,
    rules_dataframe: pd.DataFrame,
    rules_column_name: str,
    rules_label: str
) -> tuple[list[str], list[str]]:
    """
    Evaluate every rule in a rules dataframe against the drawing.

    Returns:
    - statuses list aligned to dataframe rows
    - remarks list aligned to dataframe rows
    """
    (
        non_empty_rule_texts,
        original_row_indexes,
        statuses,
        remarks
    ) = prepare_non_empty_rules(rules_dataframe, rules_column_name)

    non_empty_rule_count = len(non_empty_rule_texts)
    total_batches = ceil(non_empty_rule_count / RULES_PER_BATCH) if non_empty_rule_count > 0 else 0

    print(
        f"[INFO] [{rules_label}] "
        f"{non_empty_rule_count} non-empty rules found. "
        f"Processing in {total_batches} batch call(s)."
    )

    for batch_number, (rules_batch, batch_start_index) in enumerate(
        split_list_into_batches(non_empty_rule_texts, RULES_PER_BATCH),
        start=1
    ):
        batch_rule_start = batch_start_index + 1
        batch_rule_end = batch_start_index + len(rules_batch)

        print(
            f"[INFO] [{rules_label}] "
            f"Processing batch {batch_number}/{total_batches} "
            f"(rules {batch_rule_start} to {batch_rule_end})"
        )

        try:
            raw_response_text = evaluate_rules_batch_with_gemini(
                gemini_client=gemini_client,
                drawing_image_parts=drawing_image_parts,
                rules_batch=rules_batch
            )

            cleaned_response_text = remove_markdown_code_fences(raw_response_text)
            parsed_response = json.loads(cleaned_response_text)

            if not isinstance(parsed_response, list):
                raise ValueError(
                    f"Expected Gemini response to be a list, but got {type(parsed_response).__name__}."
                )

            if len(parsed_response) != len(rules_batch):
                raise ValueError(
                    f"Expected {len(rules_batch)} result objects, but got {len(parsed_response)}."
                )

            for batch_item_index, result_item in enumerate(parsed_response):
                original_row_index = original_row_indexes[batch_start_index + batch_item_index]

                if isinstance(result_item, dict):
                    statuses[original_row_index] = result_item.get("status", "Error")
                    remarks[original_row_index] = result_item.get("remarks", "")
                else:
                    statuses[original_row_index] = "Error"
                    remarks[original_row_index] = "Invalid response item format"

            print(f"[INFO] [{rules_label}] Batch {batch_number} completed successfully.")

        except Exception as error:
            print(f"[ERROR] [{rules_label}] Batch {batch_number} failed: {error}")

            for batch_item_index in range(len(rules_batch)):
                original_row_index = original_row_indexes[batch_start_index + batch_item_index]
                statuses[original_row_index] = "Error"
                remarks[original_row_index] = f"Error: {error}"

        if batch_number < total_batches:
            print(f"[INFO] Waiting {API_CALL_DELAY_SECONDS} seconds before next API call...")
            time.sleep(API_CALL_DELAY_SECONDS)

    return statuses, remarks

# Output building

def build_output_section(
    rules_dataframe: pd.DataFrame,
    rules_column_name: str,
    statuses: list[str],
    remarks: list[str]
) -> pd.DataFrame:
    """
    Build one output dataframe section with Rule, Status, and Remarks columns.
    """
    section_dataframe = pd.DataFrame({
        "Rule": rules_dataframe[rules_column_name].values,
        "Status": statuses,
        "Remarks": remarks,
    })

    return section_dataframe


def build_final_output_dataframe(
    common_rules_dataframe: pd.DataFrame, common_rules_column_name: str,
    common_statuses: list[str],
    common_remarks: list[str],
    drawing_rules_dataframe: pd.DataFrame,
    drawing_rules_column_name: str,
    drawing_statuses: list[str],
    drawing_remarks: list[str]) -> tuple[pd.DataFrame, list[int]]:
    """
    Build the final Excel output dataframe with:
    - a titled common rules section
    - a titled drawing-specific rules section

    Returns:
    - final output dataframe
    - list of Excel row numbers for section title rows
    """
    common_rules_section = build_output_section(
        rules_dataframe=common_rules_dataframe,
        rules_column_name=common_rules_column_name,
        statuses=common_statuses,
        remarks=common_remarks
    )

    drawing_rules_section = build_output_section(
        rules_dataframe=drawing_rules_dataframe,
        rules_column_name=drawing_rules_column_name,
        statuses=drawing_statuses,
        remarks=drawing_remarks
    )

    common_section_title_row = pd.DataFrame([
        {"Rule": "--- Common Rules ---", "Status": "", "Remarks": ""}
    ])

    drawing_section_title_row = pd.DataFrame([
        {"Rule": "--- Drawing-Specific Rules ---", "Status": "", "Remarks": ""}
    ])

    final_output_dataframe = pd.concat(
        [
            common_section_title_row,
            common_rules_section,
            drawing_section_title_row,
            drawing_rules_section
        ],
        ignore_index=True
    )

    # Excel row numbering is 1-based, and row 1 is reserved for column headers.
    common_section_title_excel_row = 2
    drawing_section_title_excel_row = len(common_rules_section) + 3

    section_title_excel_rows = [
        common_section_title_excel_row,
        drawing_section_title_excel_row
    ]

    return final_output_dataframe, section_title_excel_rows


def apply_excel_formatting(output_excel_path: str, section_title_excel_rows: list[int]) -> None:
    """
    Open the saved Excel file and apply formatting to:
    - the section title rows
    - the top column header row
    """
    workbook = load_workbook(output_excel_path)
    worksheet = workbook.active

    section_title_font = Font(bold=True, size=15)
    section_title_fill = PatternFill(fill_type="solid", fgColor="D9E1F2")
    section_title_alignment = Alignment(horizontal="left", vertical="center")

    for row_number in section_title_excel_rows:
        for cell in worksheet[row_number]:
            cell.font = section_title_font
            cell.fill = section_title_fill
            cell.alignment = section_title_alignment

    column_header_font = Font(bold=True, size=11)
    column_header_fill = PatternFill(fill_type="solid", fgColor="BDD7EE")
    column_header_alignment = Alignment(horizontal="center", vertical="center")

    for cell in worksheet[1]:
        cell.font = column_header_font
        cell.fill = column_header_fill
        cell.alignment = column_header_alignment

    workbook.save(output_excel_path)

# Main QC process

def run_quality_check_process(
    drawing_pdf_path: str,
    drawing_rules_excel_path: str,
    output_excel_path: str
) -> None:
    """
    Run the full quality check process:
    1. Convert the drawing PDF into images
    2. Load common rules
    3. Load drawing-specific rules
    4. Evaluate both rule sets against the drawing
    5. Save a combined Excel output
    """
    print("=== QC Process Started ===")
    print(f"[INFO] Drawing PDF path: {drawing_pdf_path}")
    print(f"[INFO] Drawing-specific rules path: {drawing_rules_excel_path}")
    print(f"[INFO] Common rules path: {COMMON_RULES_FILE_PATH}")
    print(f"[INFO] Output Excel path: {output_excel_path}")

    drawing_image_parts = convert_pdf_to_gemini_image_parts(drawing_pdf_path)
    gemini_client = initialize_gemini_client()

    common_rules_dataframe, common_rules_column_name = load_rules_excel_file(
        COMMON_RULES_FILE_PATH,
        "common"
    )

    common_statuses, common_remarks = evaluate_rules_against_drawing(
        gemini_client=gemini_client,
        drawing_image_parts=drawing_image_parts,
        rules_dataframe=common_rules_dataframe,
        rules_column_name=common_rules_column_name,
        rules_label="Common Rules"
    )

    drawing_rules_dataframe, drawing_rules_column_name = load_rules_excel_file(
        drawing_rules_excel_path,
        "drawing-specific"
    )

    drawing_statuses, drawing_remarks = evaluate_rules_against_drawing(
        gemini_client=gemini_client,
        drawing_image_parts=drawing_image_parts,
        rules_dataframe=drawing_rules_dataframe,
        rules_column_name=drawing_rules_column_name,
        rules_label="Drawing-Specific Rules"
    )

    final_output_dataframe, section_title_excel_rows = build_final_output_dataframe(
        common_rules_dataframe=common_rules_dataframe,
        common_rules_column_name=common_rules_column_name,
        common_statuses=common_statuses,
        common_remarks=common_remarks,
        drawing_rules_dataframe=drawing_rules_dataframe,
        drawing_rules_column_name=drawing_rules_column_name,
        drawing_statuses=drawing_statuses,
        drawing_remarks=drawing_remarks,
    )

    output_directory = os.path.dirname(os.path.abspath(output_excel_path))
    os.makedirs(output_directory, exist_ok=True)

    final_output_dataframe.to_excel(output_excel_path, index=False)
    apply_excel_formatting(output_excel_path, section_title_excel_rows)

    print(f"[INFO] QC result file saved successfully: {output_excel_path}")
    print("=== QC Process Completed ===")

# Script entry point

if __name__ == "__main__":
    try:
        drawing_pdf_path = input("Enter drawing PDF path: ").strip()
        drawing_rules_excel_path = input("Enter drawing-specific rules Excel path: ").strip()

        output_directory = os.path.join(
            os.path.dirname(os.path.abspath(__file__)),
            "output"
        )
        os.makedirs(output_directory, exist_ok=True)

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_excel_path = os.path.join(
            output_directory,
            f"qc_results_{timestamp}.xlsx"
        )

        run_quality_check_process(
            drawing_pdf_path=drawing_pdf_path,
            drawing_rules_excel_path=drawing_rules_excel_path,
            output_excel_path=output_excel_path
        )

    except Exception as error:
        print("### FATAL ERROR ###")
        print(str(error))
        print("### END ERROR ###")